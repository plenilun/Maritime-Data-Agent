from __future__ import annotations

import io
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from . import agent_memory, code_dictionary, db, relationships
from .agent import AgentDatasetNotFound, run_data_agent
from .agent_eval import run_agent_evals
from .agent_planner import DEFAULT_PLANNER
from .agent_tools import DEFAULT_TOOL_REGISTRY
from .llm import ModelConfig
from .multitable import JOIN_KEY_PAIRS, TABLE_PURPOSES

STATIC_DIR = Path(__file__).resolve().parent / "static"


@asynccontextmanager
async def lifespan(_: FastAPI):
    db.init_db()
    agent_memory.init_memory()
    code_dictionary.initialize_default_dictionary()
    yield


app = FastAPI(title="智能问数", version="1.0.0", lifespan=lifespan)


class TermCreate(BaseModel):
    term: str = Field(min_length=1, max_length=100)
    definition: str = Field(min_length=1, max_length=1000)
    synonyms: str = Field(default="", max_length=500)
    dataset_id: str | None = None


class CodeEntryUpdate(BaseModel):
    description: str = Field(min_length=1, max_length=500)
    synonyms: str = Field(default="", max_length=1000)


class CodeEntryCreate(CodeEntryUpdate):
    code_type: str = Field(min_length=1, max_length=100)
    code_value: str = Field(min_length=1, max_length=100)


class CodeBindingCreate(BaseModel):
    dataset_id: str
    table_name: str
    column_name: str
    code_type: str
    enabled: bool = True


class RelationshipCreate(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    left_dataset_id: str
    left_field: str = Field(min_length=1, max_length=200)
    right_dataset_id: str
    right_field: str = Field(min_length=1, max_length=200)
    meaning: str = Field(min_length=1, max_length=1000)
    left_grain: str = Field(min_length=1, max_length=500)
    right_grain: str = Field(min_length=1, max_length=500)
    enabled: bool = True


class RelationshipStatus(BaseModel):
    enabled: bool


class AskRequest(BaseModel):
    dataset_id: str | None = None
    secondary_dataset_id: str | None = None
    dataset_ids: list[str] = Field(default_factory=list)
    question: str = Field(min_length=1, max_length=2000)
    model: dict[str, Any]


class DatasetRename(BaseModel):
    name: str = Field(min_length=1, max_length=100)


def clean_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.dropna(how="all").copy()
    frame.columns = [str(col).strip() or f"column_{idx + 1}" for idx, col in enumerate(frame.columns)]
    seen: dict[str, int] = {}
    names: list[str] = []
    for name in frame.columns:
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    frame.columns = names
    for col in frame.columns:
        if pd.api.types.is_datetime64_any_dtype(frame[col]):
            frame[col] = frame[col].dt.strftime("%Y-%m-%d %H:%M:%S")
    return frame.where(pd.notna(frame), None)


def sqlite_type(series: pd.Series) -> str:
    if pd.api.types.is_integer_dtype(series):
        return "INTEGER"
    if pd.api.types.is_float_dtype(series):
        return "REAL"
    return "TEXT"


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/agent/manifest")
def agent_manifest() -> dict[str, Any]:
    return {
        "agent": "MaritimeDataAgent",
        "planner": DEFAULT_PLANNER.describe(),
        "tools": DEFAULT_TOOL_REGISTRY.describe(),
    }


@app.get("/api/agent/memory")
def agent_memory_list(limit: int = Query(default=20, ge=1, le=100)) -> list[dict[str, Any]]:
    return agent_memory.list_memories(limit)


@app.get("/api/agent/evals")
def agent_evals() -> dict[str, Any]:
    return run_agent_evals()


@app.get("/api/datasets")
def datasets() -> list[dict[str, Any]]:
    return db.list_datasets()


@app.post("/api/datasets")
async def upload_dataset(
    file: UploadFile = File(...),
    name: str = Form(default=""),
    sheet_name: str = Form(default=""),
) -> dict[str, Any]:
    filename = file.filename or "dataset"
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx", ".xls"}:
        raise HTTPException(400, "仅支持 CSV、XLSX 和 XLS 文件")
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(400, "文件不能超过 50 MB")
    try:
        if suffix == ".csv":
            try:
                frame = pd.read_csv(io.BytesIO(content), encoding="utf-8-sig")
            except UnicodeDecodeError:
                frame = pd.read_csv(io.BytesIO(content), encoding="gb18030")
        else:
            frame = pd.read_excel(io.BytesIO(content), sheet_name=sheet_name or 0)
    except Exception as exc:
        raise HTTPException(400, f"无法读取数据文件：{exc}") from exc
    frame = clean_frame(frame)
    if not len(frame.columns):
        raise HTTPException(400, "文件中没有可用字段")
    dataset_id = uuid.uuid4().hex
    display_name = name.strip() or Path(filename).stem
    table_name = db.safe_identifier(display_name)
    try:
        with db.connection() as conn:
            frame.to_sql(table_name, conn, index=False, if_exists="fail")
        columns = [{"name": col, "type": sqlite_type(frame[col])} for col in frame.columns]
        db.save_dataset(dataset_id, display_name, table_name, filename, len(frame), columns)
        code_dictionary.sync_default_bindings()
    except Exception as exc:
        with db.connection() as conn:
            conn.execute(f'DROP TABLE IF EXISTS "{table_name}"')
        raise HTTPException(500, f"导入数据失败：{exc}") from exc
    return db.get_dataset(dataset_id) or {}


@app.delete("/api/datasets/{dataset_id}")
def remove_dataset(dataset_id: str) -> dict[str, bool]:
    if not db.delete_dataset(dataset_id):
        raise HTTPException(404, "数据集不存在")
    from .llm import clear_sql_cache
    clear_sql_cache()
    return {"ok": True}


@app.patch("/api/datasets/{dataset_id}")
def rename_dataset(dataset_id: str, payload: DatasetRename) -> dict[str, Any]:
    try:
        dataset = db.rename_dataset(dataset_id, payload.name)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    if not dataset:
        raise HTTPException(404, "数据表不存在")
    from .llm import clear_sql_cache
    clear_sql_cache()
    return dataset


@app.get("/api/relationships")
def relationship_list() -> list[dict[str, Any]]:
    return relationships.list_relationships()


@app.get("/api/relationship-rules")
def relationship_rules() -> dict[str, Any]:
    return {"join_key_pairs": JOIN_KEY_PAIRS, "table_purposes": TABLE_PURPOSES}


# 暂时注释掉可配置关系的API，待multitable.py实现后启用
# @app.post("/api/relationships")
# def create_relationship(payload: RelationshipCreate) -> dict[str, Any]:
#     try:
#         return relationships.save_relationship(payload.model_dump())
#     except ValueError as exc:
#         raise HTTPException(400, str(exc)) from exc


# @app.put("/api/relationships/{relationship_id}")
# def edit_relationship(relationship_id: str, payload: RelationshipCreate) -> dict[str, Any]:
#     try:
#         return relationships.save_relationship(payload.model_dump(), relationship_id)
#     except LookupError as exc:
#         raise HTTPException(404, str(exc)) from exc
#     except ValueError as exc:
#         raise HTTPException(400, str(exc)) from exc


# @app.patch("/api/relationships/{relationship_id}/status")
# def relationship_status(relationship_id: str, payload: RelationshipStatus) -> dict[str, Any]:
#     try:
#         item = relationships.set_enabled(relationship_id, payload.enabled)
#     except ValueError as exc:
#         raise HTTPException(400, str(exc)) from exc
#     if not item:
#         raise HTTPException(404, "关系不存在")
#     return item


# @app.delete("/api/relationships/{relationship_id}")
# def remove_relationship(relationship_id: str) -> dict[str, bool]:
#     if not relationships.delete_relationship(relationship_id):
#         raise HTTPException(404, "关系不存在")
#     return {"ok": True}


# 暫時注釋掉檢查關係的API，待multitable.py實現後啟用
# @app.get("/api/relationships/{relationship_id}/check")
# def check_relationship(relationship_id: str) -> dict[str, Any]:
#     try:
#         return relationships.inspect_relationship(relationship_id)
#     except LookupError as exc:
#         raise HTTPException(404, str(exc)) from exc
#     except ValueError as exc:
#         raise HTTPException(400, str(exc)) from exc


@app.get("/api/datasets/{dataset_id}/preview")
def dataset_preview(dataset_id: str, limit: int = Query(default=100, ge=1, le=200)) -> dict[str, Any]:
    preview = db.preview_dataset(dataset_id, limit)
    if not preview:
        raise HTTPException(404, "数据集不存在")
    return preview


@app.get("/api/admin/code-versions")
def code_versions() -> list[dict[str, Any]]:
    return code_dictionary.list_versions()


@app.post("/api/admin/code-versions/{version_id}/activate")
def activate_code_version(version_id: str) -> dict[str, Any]:
    version = code_dictionary.activate_version(version_id)
    if not version:
        raise HTTPException(404, "编码字典版本不存在")
    return version


@app.post("/api/admin/code-import")
async def import_code_dictionary(file: UploadFile = File(...), dry_run: bool = Query(default=False)) -> dict[str, Any]:
    filename = file.filename or "fm_code.xlsx"
    if Path(filename).suffix.lower() not in {".xlsx", ".xls"}:
        raise HTTPException(400, "仅支持 XLSX 和 XLS 文件")
    content = await file.read()
    if not content:
        raise HTTPException(400, "上传的文件为空")
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "编码文件不能超过 10 MB")
    preview = code_dictionary.parse_workbook(content)
    if dry_run:
        return {
            "valid": not preview["errors"],
            "entry_count": len(preview["entries"]),
            "errors": preview["errors"],
            "duplicates": preview["duplicates"],
            "preview": preview["entries"][:20],
        }
    try:
        return code_dictionary.import_workbook(content, filename, activate=True)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.get("/api/admin/code-entries")
def code_entries(q: str = Query(default="", max_length=100), code_type: str = Query(default="", max_length=100)) -> list[dict[str, Any]]:
    return code_dictionary.list_entries(q, code_type)


@app.post("/api/admin/code-entries")
def create_code_entry(payload: CodeEntryCreate) -> dict[str, Any]:
    try:
        return code_dictionary.create_entry(
            payload.code_type, payload.code_value, payload.description, payload.synonyms
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.put("/api/admin/code-entries/{entry_id}")
def edit_code_entry(entry_id: str, payload: CodeEntryUpdate) -> dict[str, Any]:
    entry = code_dictionary.update_entry(entry_id, payload.description, payload.synonyms)
    if not entry:
        raise HTTPException(404, "编码项不存在")
    return entry


@app.delete("/api/admin/code-entries/{entry_id}")
def remove_code_entry(entry_id: str) -> dict[str, bool]:
    if not code_dictionary.delete_entry(entry_id):
        raise HTTPException(404, "编码项不存在")
    return {"ok": True}


@app.get("/api/admin/code-bindings")
def code_bindings(dataset_id: str = Query(default="", max_length=100)) -> list[dict[str, Any]]:
    return code_dictionary.list_bindings(dataset_id)


@app.post("/api/admin/code-bindings")
def create_code_binding(payload: CodeBindingCreate) -> dict[str, Any]:
    try:
        return code_dictionary.save_binding(
            payload.dataset_id, payload.table_name, payload.column_name, payload.code_type, payload.enabled
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.delete("/api/admin/code-bindings/{binding_id}")
def remove_code_binding(binding_id: str) -> dict[str, bool]:
    if not code_dictionary.delete_binding(binding_id):
        raise HTTPException(404, "字段绑定不存在")
    return {"ok": True}


@app.get("/api/terms")
def terms(dataset_id: str | None = None, q: str = Query(default="", max_length=100)) -> list[dict[str, Any]]:
    return db.list_terms(dataset_id, q)


@app.post("/api/terms")
def create_term(payload: TermCreate) -> dict[str, Any]:
    if payload.dataset_id and not db.get_dataset(payload.dataset_id):
        raise HTTPException(404, "关联的数据集不存在")
    return db.add_term(payload.term, payload.definition, payload.synonyms, payload.dataset_id)


@app.put("/api/terms/{term_id}")
def edit_term(term_id: str, payload: TermCreate) -> dict[str, Any]:
    if payload.dataset_id and not db.get_dataset(payload.dataset_id):
        raise HTTPException(404, "关联的数据集不存在")
    term = db.update_term(term_id, payload.term, payload.definition, payload.synonyms, payload.dataset_id)
    if not term:
        raise HTTPException(404, "术语不存在")
    return term


@app.post("/api/terms/import")
async def import_terms(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = file.filename or "terms.xlsx"
    if Path(filename).suffix.lower() not in {".xlsx", ".xls"}:
        raise HTTPException(400, "仅支持 XLSX 和 XLS 文件")
    content = await file.read()
    if not content:
        raise HTTPException(400, "上传的文件为空")
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "术语文件不能超过 10 MB")
    try:
        frame = pd.read_excel(io.BytesIO(content), dtype=object)
    except Exception as exc:
        raise HTTPException(400, f"无法读取术语文件：{exc}") from exc

    frame.columns = [str(column).strip() for column in frame.columns]
    columns = ["术语", "定义", "同义词", "关联数据表"]
    missing = set(columns).difference(frame.columns)
    if missing:
        raise HTTPException(400, f"缺少字段：{'、'.join(sorted(missing))}")

    datasets_by_name: dict[str, list[dict[str, Any]]] = {}
    for dataset in db.list_datasets():
        datasets_by_name.setdefault(dataset["name"].strip(), []).append(dataset)
    existing = {(item["term"].strip(), item.get("dataset_id")) for item in db.list_terms()}
    pending: set[tuple[str, str | None]] = set()
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    skipped = 0

    def cell_text(value: Any) -> str:
        return "" if pd.isna(value) else str(value).strip()

    for index, record in frame[columns].iterrows():
        row_number = index + 2
        term = cell_text(record["术语"])
        definition = cell_text(record["定义"])
        synonyms = cell_text(record["同义词"])
        dataset_name = cell_text(record["关联数据表"])
        is_global = dataset_name in {"全局术语", "全局"}
        if not any((term, definition, synonyms, dataset_name)):
            continue
        if not term:
            errors.append(f"第 {row_number} 行：术语不能为空")
        if not definition:
            errors.append(f"第 {row_number} 行：定义不能为空")
        if len(term) > 100:
            errors.append(f"第 {row_number} 行：术语不能超过 100 个字符")
        if len(definition) > 1000:
            errors.append(f"第 {row_number} 行：定义不能超过 1000 个字符")
        if len(synonyms) > 500:
            errors.append(f"第 {row_number} 行：同义词不能超过 500 个字符")
        dataset_id = None
        if dataset_name and not is_global:
            matches = datasets_by_name.get(dataset_name, [])
            if not matches:
                errors.append(f"第 {row_number} 行：关联数据表“{dataset_name}”不存在")
            elif len(matches) > 1:
                errors.append(f"第 {row_number} 行：关联数据表“{dataset_name}”名称不唯一")
            else:
                dataset_id = matches[0]["id"]
        if not term or not definition or (dataset_name and not is_global and dataset_id is None):
            continue
        key = (term, dataset_id)
        if key in existing or key in pending:
            skipped += 1
            continue
        pending.add(key)
        rows.append({"term": term, "definition": definition, "synonyms": synonyms, "dataset_id": dataset_id})

    if errors:
        detail = "；".join(errors[:20])
        if len(errors) > 20:
            detail += f"；另有 {len(errors) - 20} 个错误"
        raise HTTPException(400, detail)
    if not rows and not skipped:
        raise HTTPException(400, "文件中没有可导入的术语")
    db.add_terms(rows)
    return {"imported": len(rows), "skipped": skipped, "total": len(rows) + skipped}


def build_terms_export() -> io.BytesIO:
    """Build a round-trip compatible terms workbook in memory."""
    datasets = {item["id"]: item["name"] for item in db.list_datasets()}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "术语库"
    headers = ["术语", "定义", "同义词", "关联数据表"]
    sheet.append(headers)
    for item in db.list_terms():
        dataset_id = item.get("dataset_id")
        sheet.append([
            item["term"],
            item["definition"],
            item.get("synonyms", ""),
            datasets.get(dataset_id, "全局术语" if not dataset_id else ""),
        ])

    header_fill = PatternFill("solid", fgColor="635BFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(1, sheet.max_row)}"
    for column, width in {"A": 24, "B": 64, "C": 36, "D": 30}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@app.get("/api/terms/export")
def export_terms() -> StreamingResponse:
    filename = f"terms_export_{datetime.now(timezone.utc):%Y%m%d_%H%M%S}.xlsx"
    return StreamingResponse(
        build_terms_export(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.delete("/api/terms/{term_id}")
def remove_term(term_id: str) -> dict[str, bool]:
    if not db.delete_term(term_id):
        raise HTTPException(404, "术语不存在")
    return {"ok": True}


@app.post("/api/model/test")
async def test_model(payload: dict[str, Any]) -> dict[str, str]:
    from .llm import chat

    try:
        config = ModelConfig.from_payload(payload)
        answer = await chat(config, [{"role": "user", "content": "只回复：连接成功"}])
        return {"message": answer.content.strip()}
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(400, str(exc)) from exc


@app.post("/api/ask")
async def ask(payload: AskRequest) -> dict[str, Any]:
    try:
        return await run_data_agent(
            question=payload.question,
            requested_dataset_id=payload.dataset_id,
            model_payload=payload.model,
            dataset_ids=payload.dataset_ids,
            secondary_dataset_id=payload.secondary_dataset_id,
        )
    except AgentDatasetNotFound as exc:
        raise HTTPException(404, str(exc)) from exc
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(400, str(exc)) from exc


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/admin/code-dictionary", include_in_schema=False)
def code_dictionary_admin() -> FileResponse:
    return FileResponse(STATIC_DIR / "code_admin.html")


@app.get("/{path:path}", include_in_schema=False)
def frontend(path: str = "") -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")
